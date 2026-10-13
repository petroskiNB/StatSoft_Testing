"""Проверки XLSX"""

from contextlib import closing
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from openpyxl import Workbook
from pandas.testing import assert_frame_equal

from table_import import SourceReadError, XlsxSource


class XlsxSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "sales.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Продажи"
        sheet.append(["id", "amount"])
        for number in range(1, 11):
            sheet.append([number, None if number in (2, 6) else number * 10])
        other = book.create_sheet("Справочник")
        other.append(["name"])
        other.append(["Москва"])
        book.create_sheet("Пустой")
        book.create_sheet("Заголовки").append(["id", "amount"])
        book.save(self.path)
        book.close()
        self.source = XlsxSource(self.path)

    def test_list_sheets_and_select_sheet(self):
        self.assertEqual(
            self.source.list_sheets(),
            ["Продажи", "Справочник", "Пустой", "Заголовки"],
        )
        result = pd.concat(list(self.source.iter_chunks("Справочник")), ignore_index=True)
        self.assertEqual(result["name"].tolist(), ["Москва"])

    def test_scan_is_limited(self):
        metadata = self.source.scan("Продажи", sample_rows=3)
        self.assertEqual(metadata.name, "Продажи")
        self.assertEqual(metadata.sampled_rows, 3)
        self.assertEqual(metadata.column_count, 2)
        self.assertIsNone(metadata.row_count)
        self.assertEqual(metadata.columns[0].dtype, "int64")
        self.assertEqual(metadata.columns[1].missing_count, 1)

    def test_chunks_and_final_dataframe(self):
        expected = pd.DataFrame({
            "id": range(1, 11),
            "amount": [10.0, None, 30.0, 40.0, 50.0, None, 70.0, 80.0, 90.0, 100.0],
        })
        for size in (4, 20):
            with self.subTest(size=size):
                chunks = list(self.source.iter_chunks("Продажи", chunk_size=size))
                if size == 4:
                    self.assertEqual([len(chunk) for chunk in chunks], [4, 4, 2])
                assert_frame_equal(pd.concat(chunks, ignore_index=True), expected)

    def test_continue_and_close(self):
        with closing(self.source.iter_chunks("Продажи", chunk_size=4)) as chunks:
            self.assertEqual(next(chunks)["id"].tolist(), [1, 2, 3, 4])
            self.assertEqual(next(chunks)["id"].tolist(), [5, 6, 7, 8])
        self.path.unlink()

    def test_empty_and_header_only(self):
        for name, column_count in (("Пустой", 0), ("Заголовки", 2)):
            with self.subTest(sheet=name):
                metadata = self.source.scan(name)
                self.assertEqual(metadata.sampled_rows, 0)
                self.assertEqual(metadata.column_count, column_count)
                chunks = list(self.source.iter_chunks(name))
                self.assertEqual(len(chunks), 1)
                self.assertTrue(chunks[0].empty)

    def test_missing_sheet(self):
        with self.assertRaises(SourceReadError) as caught:
            self.source.scan("Нет такого листа")
        self.assertIn("Нет такого листа", str(caught.exception))
        self.assertIn("Продажи", str(caught.exception))

    def test_missing_and_corrupt_file(self):
        self.path.unlink()
        with self.assertRaises(SourceReadError) as caught:
            self.source.list_sheets()
        self.assertIsInstance(caught.exception.__cause__, FileNotFoundError)
        self.path.write_bytes(b"not an xlsx archive")
        with self.assertRaises(SourceReadError):
            self.source.scan("Продажи")

    def test_invalid_sizes(self):
        for size in (0, -1, True, 1.5):
            with self.subTest(size=size):
                with self.assertRaises(ValueError):
                    self.source.scan("Продажи", sample_rows=size)
                with self.assertRaises(ValueError):
                    list(self.source.iter_chunks("Продажи", chunk_size=size))

    def test_invalid_headers_and_late_wide_row(self):
        for header in (["id", "id"], ["id", None, "amount"]):
            with self.subTest(header=header):
                book = Workbook()
                book.active.append(header)
                book.save(self.path)
                book.close()
                with self.assertRaises(SourceReadError):
                    self.source.scan("Sheet")
        book = Workbook()
        book.active.append(["id"])
        book.active.append([1])
        book.active.append([2, "extra"])
        book.save(self.path)
        book.close()
        # Ошибка за границей выборки не должна мешать сканированию
        self.assertEqual(self.source.scan("Sheet", sample_rows=1).sampled_rows, 1)
        with self.assertRaises(SourceReadError):
            list(self.source.iter_chunks("Sheet", chunk_size=1))


if __name__ == "__main__":
    unittest.main()
