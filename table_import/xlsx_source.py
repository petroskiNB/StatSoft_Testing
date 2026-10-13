""" Потоковое чтение листов XLSX """

from collections.abc import Generator
from contextlib import closing, contextmanager
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from .exceptions import SourceReadError
from .models import ColumnMetadata, SheetMetadata


class XlsxSource:
    """Первая строка листа - заголовки, а остальные - данные"""

    def __init__(self, path: str | Path):
        self.path = Path(path)

    @staticmethod
    def _validate_size(value: int, name: str) -> None:
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            raise ValueError(f"{name} должен быть положительным целым числом")

    @contextmanager
    def _open_workbook(self, operation: str):
        """Закрывает книгу и файл при успехе, ошибке и закрытии генератора"""
        try:
            with self.path.open("rb") as file:
                workbook = load_workbook(file, read_only=True, data_only=True)
                try:
                    yield workbook
                finally:
                    workbook.close()
        except (
            OSError, ValueError, KeyError, BadZipFile,
            InvalidFileException, ParseError, EOFError
        ) as exc:
            raise SourceReadError(self.path, operation, str(exc)) from exc

    def list_sheets(self) -> list[str]:
        """Возвращает названия обычных листов в порядке книги."""
        with self._open_workbook("получение списка листов XLSX") as workbook:
            return [worksheet.title for worksheet in workbook.worksheets]

    def scan(self, sheet: str, sample_rows: int = 100) -> SheetMetadata:
        """Определяет типы и пропуски только по ограниченной выборке"""
        self._validate_size(sample_rows, "sample_rows")
        # Берем одну порцию, затем сразу освобождаем ресурсы генератора
        with closing(self.iter_chunks(sheet, chunk_size=sample_rows)) as chunks:
            sample = next(chunks)

        columns = tuple(
            ColumnMetadata(
                name=name,
                dtype=str(sample[name].dtype),
                missing_count=int(sample[name].isna().sum())
            )
            for name in sample.columns
        )
        return SheetMetadata(
            name=sheet,
            row_count=None,
            column_count=len(sample.columns),
            sampled_rows=len(sample),
            columns=columns,
        )

    def iter_chunks(
        self, sheet: str, chunk_size: int = 1000,
    ) -> Generator[pd.DataFrame, None, None]:
        """Выдает порции. Пустой лист даёт один пустой DataFrame"""
        self._validate_size(chunk_size, "chunk_size")
        if not isinstance(sheet, str) or not sheet:
            raise ValueError("sheet должен быть непустой строкой")

        with self._open_workbook(f"чтение листа XLSX '{sheet}'") as workbook:
            names = [worksheet.title for worksheet in workbook.worksheets]
            if sheet not in names:
                raise SourceReadError(
                    self.path, "выбор листа XLSX",
                    f"Лист '{sheet}' не найден. Доступные листы: {names}",
                )
            worksheet = workbook[sheet]
            # Размеры, сохраненные создателем XLSX, бывают неточными
            # Читаем реальные строки, не вычисляя полные размеры листа
            worksheet.reset_dimensions()
            with closing(worksheet.iter_rows(values_only=True)) as rows:
                first_row = next(rows, None)
                if first_row is None:
                    yield pd.DataFrame()
                    return
                header = list(first_row)
                while header and header[-1] is None:
                    header.pop()
                if not header:
                    raise ValueError("Первая строка должна содержать заголовки")

                if any(value is None or not str(value).strip() for value in header):
                    raise ValueError("Названия столбцов должны быть непустыми")
                names = [str(value) for value in header]
                if len(set(names)) != len(names):
                    raise ValueError("Названия столбцов должны быть уникальными")

                batch = []
                emitted = False
                for row_number, row in enumerate(rows, start=2):
                    if any(value is not None for value in row[len(names):]):
                        raise ValueError(
                            f"Строка {row_number}: данные выходят за пределы заголовков"
                        )
                    # Короткие строки дополняем пропусками до ширины таблицы
                    values = list(row[:len(names)])
                    values.extend([None] * (len(names) - len(values)))
                    batch.append(values)
                    if len(batch) == chunk_size:
                        yield pd.DataFrame(batch, columns=names)
                        emitted = True
                        batch = []

                if batch or not emitted:
                    yield pd.DataFrame(batch, columns=names)
